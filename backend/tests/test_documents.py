"""知识库管理回归：内存 SQL、临时私有目录、假向量服务，不读取真实 Key。"""

from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from concurrent.futures import ThreadPoolExecutor
import threading
import unittest
from unittest.mock import patch
import uuid

import httpx
from fastapi import FastAPI, HTTPException
from fastapi.exceptions import RequestValidationError
from pydantic_settings import BaseSettings
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool


_settings_init = BaseSettings.__init__


def _offline_settings(self, *args, **kwargs):
    kwargs["_env_file"] = None
    _settings_init(self, *args, **kwargs)


with patch.dict(os.environ, {
    "DATABASE_URL": "sqlite://", "JWT_SECRET_KEY": uuid.uuid4().hex,
    "LLM_API_KEY": "", "EMBEDDING_API_KEY": "",
}), patch.object(BaseSettings, "__init__", _offline_settings):
    from app.api.document import router
    from app.common.exceptions import (
        BusinessException, business_exception_handler, http_excpetion_hadler,
        validation_excpetion_hadler,
    )
    from app.database import Base, get_db
    from app.models.document import Document
    from app.models.user import User
    from app.services import document_service, rag_vector_service
    from app.utils.jwt import create_access_token


class FakeVectors:
    def __init__(self):
        self.rows = {}
        self.deleted = []
        self.fail_index_filename = None
        self.fail_delete = False

    def index_document(self, document_id, filename, chunks):
        self.rows[document_id] = {"filename": filename, "chunks": chunks}
        if filename == self.fail_index_filename:
            # 真实库可能写到一半才失败：必须连失败这份也做补偿。
            self.rows[document_id]["chunks"] = chunks[:1]
            raise BusinessException("向量写入异常", code=502)

    def delete_document(self, document_id):
        self.deleted.append(document_id)
        if self.fail_delete:
            raise BusinessException("向量清理异常", code=502)
        self.rows.pop(document_id, None)


def docx_bytes():
    from docx import Document as WordDocument

    document = WordDocument()
    document.add_paragraph("DOCX 教务规定")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "请提前十分钟到教室"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(*, blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if not blank:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        })
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 30 220 Td (PDF campus rules) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


class DocumentApiTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = (Path(self.temp.name) / "private_documents").resolve()
        self.engine = create_engine(
            "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
        )
        self.addCleanup(self.engine.dispose)
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, autoflush=False)
        with self.sessions() as db:
            admin = User(username="admin", password="test-only", name="管理员", role="admin")
            student = User(username="aaa", password="test-only", name="学生", role="student")
            db.add_all([admin, student])
            db.commit()
            self.admin_id = admin.id
            self.tokens = {
                "admin": create_access_token(admin.id, admin.password),
                "student": create_access_token(student.id, student.password),
            }

        self.vectors = FakeVectors()
        for modification in (
            patch.object(document_service, "DOCUMENT_DIR", self.directory),
            patch.object(rag_vector_service, "CHROMA_DIR", Path(self.temp.name) / "vectors"),
            patch.object(document_service.rag_service, "index_document", self.vectors.index_document),
            patch.object(document_service.rag_service, "delete_document", self.vectors.delete_document),
        ):
            modification.start()
            self.addCleanup(modification.stop)
        self.app = FastAPI()
        self.app.include_router(router, prefix="/api")
        self.app.add_exception_handler(BusinessException, business_exception_handler)
        self.app.add_exception_handler(HTTPException, http_excpetion_hadler)
        self.app.add_exception_handler(RequestValidationError, validation_excpetion_hadler)

        def test_db():
            with self.sessions() as db:
                yield db

        self.app.dependency_overrides[get_db] = test_db

    async def request(self, method, path="/api/documents", *, user="admin", files=None):
        headers = {} if user is None else {"Authorization": f"Bearer {self.tokens[user]}"}
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app), base_url="http://offline"
        ) as client:
            response = await client.request(method, path, headers=headers, files=files)
        return response.status_code, response.json()

    async def upload(self, files, *, user="admin"):
        return await self.request("POST", "/api/documents/upload", user=user, files=[
            ("files", (filename, content, "application/octet-stream"))
            for filename, content in files
        ])

    async def save_one(self, filename="规定.txt", text="知识库规定"):
        status, response = await self.upload([(filename, text.encode())])
        self.assertEqual(status, 200)
        self.assertEqual(response["code"], 200, response)
        return response["data"][0]

    def rows(self):
        with self.sessions() as db:
            return db.query(Document).all()

    def files(self):
        return list(self.directory.iterdir()) if self.directory.exists() else []

    def assert_empty(self):
        self.assertEqual(self.rows(), [])
        self.assertEqual(self.files(), [])
        self.assertEqual(self.vectors.rows, {})

    async def test_four_formats_upload_private_files_and_index(self):
        status, response = await self.upload([
            ("规定.txt", "校规第一条".encode()), ("规定.md", "# 校规\n遵守上课时间".encode()),
            ("规定.pdf", pdf_bytes()), ("规定.docx", docx_bytes()),
        ])
        self.assertEqual(status, 200)
        self.assertEqual(response["code"], 200, response)
        self.assertEqual(len(response["data"]), 4)
        for item in response["data"]:
            self.assertEqual(item["uploader_id"], self.admin_id)
            self.assertEqual(item["status"], "indexed")
            self.assertGreater(item["chunk_count"], 0)
            self.assertTrue(item["create_time"])
            self.assertTrue(item["update_time"])
            self.assertNotIn("storage_path", item)
            self.assertNotIn("extracted_text", item)
        self.assertEqual(len(self.files()), 4)
        self.assertEqual(len(self.vectors.rows), 4)
        for stored in self.rows():
            self.assertEqual(Path(stored.storage_path).parent, self.directory)
            self.assertNotEqual(Path(stored.storage_path).name, stored.filename)
            self.assertTrue(stored.extracted_text.strip())

    async def test_unauthenticated_and_students_cannot_access_any_operation(self):
        row = await self.save_one()
        for user, expected_http, expected_code in ((None, 401, 401), ("student", 200, 403)):
            for method, path, files in (
                ("GET", "/api/documents", None),
                ("POST", "/api/documents/upload", [("files", ("a.txt", b"text", "text/plain"))]),
                ("GET", f"/api/documents/{row['id']}/preview", None),
                ("POST", f"/api/documents/{row['id']}/reindex", None),
                ("DELETE", f"/api/documents/{row['id']}", None),
            ):
                with self.subTest(user=user, method=method, path=path):
                    status, response = await self.request(method, path, user=user, files=files)
                    self.assertEqual(status, expected_http)
                    self.assertEqual(response["code"], expected_code)
        self.assertEqual(len(self.rows()), 1)
        self.assertEqual(len(self.files()), 1)
        self.assertEqual(len(self.vectors.rows), 1)

    async def test_scanned_pdf_rolls_back_preceding_good_file(self):
        _, response = await self.upload([
            ("正常.txt", b"valid content"), ("扫描.pdf", pdf_bytes(blank=True))
        ])
        self.assertEqual(response["code"], 400)
        self.assertIn("扫描.pdf", response["message"])
        self.assertIn("扫描版 PDF", response["message"])
        self.assertIn("OCR", response["message"])
        self.assertIn("整批已回滚", response["message"])
        self.assertEqual(len(self.vectors.deleted), 1)
        self.assert_empty()

    async def test_partial_vector_failure_compensates_failed_and_successful_documents(self):
        self.vectors.fail_index_filename = "失败.txt"
        _, response = await self.upload([
            ("正常.txt", b"valid content"), ("失败.txt", b"failed content")
        ])
        self.assertEqual(response["code"], 502)
        self.assertIn("失败.txt", response["message"])
        self.assertIn("向量写入异常", response["message"])
        self.assertEqual(len(set(self.vectors.deleted)), 2)
        self.assert_empty()

    async def test_unsupported_extension_and_bad_utf8_roll_back_entire_batch(self):
        for bad_filename, content in (("程序.exe", b"bad"), ("编码.txt", b"\xff\xfe")):
            with self.subTest(filename=bad_filename):
                _, response = await self.upload([("正常.txt", b"valid"), (bad_filename, content)])
                self.assertNotEqual(response["code"], 200)
                self.assertIn(bad_filename, response["message"])
                self.assert_empty()

    async def test_empty_files_and_excessive_sizes_fail_explicitly(self):
        _, response = await self.upload([("空白.txt", b"   \n\t")])
        self.assertEqual(response["code"], 400)
        self.assertIn("有效文本", response["message"])
        self.assert_empty()
        with patch.object(document_service, "MAX_FILE_SIZE", 8):
            _, response = await self.upload([("过大.txt", b"123456789")])
        self.assertEqual(response["code"], 400)
        self.assertIn("超过", response["message"])
        self.assert_empty()

    async def test_private_file_write_failure_rolls_back_database(self):
        with patch.object(Path, "write_bytes", side_effect=OSError("offline disk full")):
            _, response = await self.upload([("写入失败.txt", b"valid text")])
        self.assertEqual(response["code"], 500)
        self.assertIn("文件读写失败", response["message"])
        self.assert_empty()

    async def test_database_commit_failure_compensates_vectors_and_files(self):
        with patch.object(Session, "commit", side_effect=RuntimeError("offline database failed")):
            _, response = await self.upload([("数据库失败.txt", b"valid text")])
        self.assertEqual(response["code"], 500)
        self.assertIn("数据库写入失败", response["message"])
        self.assert_empty()

    async def test_cleanup_failure_is_not_reported_as_complete_rollback(self):
        self.vectors.fail_index_filename = "失败.txt"
        self.vectors.fail_delete = True
        _, response = await self.upload([("失败.txt", b"valid text")])
        self.assertEqual(response["code"], 502)
        self.assertIn("补偿清理未完成", response["message"])
        self.assertNotIn("整批已回滚", response["message"])
        self.assertEqual(self.rows(), [])
        self.assertEqual(self.files(), [])
        self.assertEqual(len(self.vectors.rows), 1)

    async def test_listing_preview_and_pagination_exclude_private_path(self):
        first = await self.save_one("先上传.md", "# 资料\n第一份")
        second = await self.save_one("后上传.txt", "第二份资料")
        _, response = await self.request("GET", "/api/documents?page=1&page_size=1")
        self.assertEqual(response["data"]["total"], 2)
        self.assertEqual(response["data"]["list"][0]["id"], second["id"])
        _, response = await self.request("GET", f"/api/documents/{first['id']}/preview")
        self.assertEqual(response["data"]["extracted_text"], "# 资料\n第一份")
        self.assertNotIn("storage_path", response["data"])

    async def test_successful_delete_removes_vectors_file_and_record(self):
        row = await self.save_one()
        _, response = await self.request("DELETE", f"/api/documents/{row['id']}")
        self.assertEqual(response["code"], 200)
        self.assertIn(row["id"], self.vectors.deleted)
        self.assert_empty()

    async def test_failed_vector_delete_preserves_record_and_original_for_retry(self):
        row = await self.save_one()
        self.vectors.fail_delete = True
        _, response = await self.request("DELETE", f"/api/documents/{row['id']}")
        self.assertEqual(response["code"], 502)
        self.assertIn("记录保留供重试", response["message"])
        self.assertEqual(self.rows()[0].status, "failed")
        self.assertEqual(len(self.files()), 1)
        self.assertEqual(len(self.vectors.rows), 1)
        self.vectors.fail_delete = False
        _, response = await self.request("DELETE", f"/api/documents/{row['id']}")
        self.assertEqual(response["code"], 200)
        self.assert_empty()

    async def test_reindex_replaces_vectors_without_duplicate_chunks(self):
        row = await self.save_one()
        stored = self.rows()[0]
        Path(stored.storage_path).write_text("新内容" * 400, encoding="utf-8")
        _, response = await self.request("POST", f"/api/documents/{row['id']}/reindex")
        self.assertEqual(response["code"], 200, response)
        self.assertEqual(response["data"]["status"], "indexed")
        self.assertGreater(response["data"]["chunk_count"], 1)
        self.assertEqual(len(self.vectors.rows), 1)
        self.assertEqual(self.rows()[0].extracted_text, "新内容" * 400)

    async def test_reindex_persists_processing_before_vector_write(self):
        row = await self.save_one()
        observed = []

        def observe_index(document_id, filename, chunks):
            with self.sessions() as db:
                observed.append(db.get(Document, document_id).status)
            self.vectors.index_document(document_id, filename, chunks)

        with patch.object(document_service.rag_service, "index_document", observe_index):
            _, response = await self.request("POST", f"/api/documents/{row['id']}/reindex")
        self.assertEqual(response["code"], 200, response)
        self.assertEqual(observed, ["processing"])
        self.assertEqual(self.rows()[0].status, "indexed")

    async def test_failed_file_delete_marks_record_for_retry_after_vectors_removed(self):
        row = await self.save_one()
        with patch.object(Path, "unlink", side_effect=OSError("offline permission failure")):
            _, response = await self.request("DELETE", f"/api/documents/{row['id']}")
        self.assertEqual(response["code"], 500)
        self.assertIn("记录保留供重试", response["message"])
        self.assertEqual(self.rows()[0].status, "failed")
        self.assertEqual(len(self.files()), 1)
        self.assertEqual(self.vectors.rows, {})
        _, response = await self.request("DELETE", f"/api/documents/{row['id']}")
        self.assertEqual(response["code"], 200)
        self.assert_empty()

    async def test_reindex_failure_retains_original_marks_failed_and_cleans_partial_vectors(self):
        row = await self.save_one("重建.txt", "旧资料")
        self.vectors.fail_index_filename = "重建.txt"
        _, response = await self.request("POST", f"/api/documents/{row['id']}/reindex")
        self.assertEqual(response["code"], 502)
        stored = self.rows()[0]
        self.assertEqual(stored.status, "failed")
        self.assertEqual(stored.chunk_count, 0)
        self.assertEqual(stored.extracted_text, "旧资料")
        self.assertIn("向量写入异常", stored.error_message)
        self.assertEqual(len(self.files()), 1)
        self.assertEqual(self.vectors.rows, {})
        self.vectors.fail_index_filename = None
        _, response = await self.request("POST", f"/api/documents/{row['id']}/reindex")
        self.assertEqual(response["code"], 200)
        self.assertEqual(response["data"]["status"], "indexed")
        self.assertIsNone(response["data"]["error_message"])

    async def test_unknown_ids_and_invalid_pagination_fail_without_mutation(self):
        for method, path in (
            ("GET", "/api/documents/9999/preview"),
            ("POST", "/api/documents/9999/reindex"),
            ("DELETE", "/api/documents/9999"),
        ):
            _, response = await self.request(method, path)
            self.assertEqual(response["code"], 404)
        for path in ("/api/documents?page=0", "/api/documents?page_size=101"):
            status, response = await self.request("GET", path)
            self.assertEqual(status, 422)
            self.assertEqual(response["code"], 422)
        self.assert_empty()

    async def test_filename_paths_cannot_control_disk_destination(self):
        row = await self.save_one("../课程规定.txt", "教务规定")
        self.assertEqual(row["filename"], "课程规定.txt")
        self.assertEqual(Path(self.rows()[0].storage_path).parent, self.directory)
        self.assertEqual(len(self.files()), 1)

    def concurrent_document(self):
        """两个真实独立 Session/连接，用文件 SQLite 避免 StaticPool 共享事务。"""
        engine = create_engine(
            f"sqlite:///{Path(self.temp.name) / 'concurrent.sqlite'}",
            connect_args={"check_same_thread": False},
        )
        self.addCleanup(engine.dispose)
        Base.metadata.create_all(engine)
        sessions = sessionmaker(bind=engine, autoflush=False)
        self.directory.mkdir(parents=True, exist_ok=True)
        path = self.directory / f"{uuid.uuid4().hex}.txt"
        path.write_text("并发重建的教务规定", encoding="utf-8")
        with sessions() as db:
            admin = User(username="concurrent-admin", password="offline", name="管理员", role="admin")
            db.add(admin)
            db.flush()
            row = Document(
                filename="并发.txt", storage_path=str(path),
                extracted_text="并发重建的教务规定", chunk_count=1,
                status="indexed", uploader_id=admin.id,
            )
            db.add(row)
            db.commit()
            document_id = row.id
        self.vectors.rows[document_id] = {"filename": "并发.txt", "chunks": ["旧分块"]}
        return sessions, document_id

    def test_failed_reindex_compensation_cannot_delete_concurrent_success(self):
        sessions, document_id = self.concurrent_document()
        first_started = threading.Event()
        second_attempted = threading.Event()
        second_finished = threading.Event()
        release_first = threading.Event()
        local = threading.local()
        trace = []

        def index(document_id, filename, chunks):
            trace.append((local.name, "index"))
            if local.name == "first":
                self.vectors.rows[document_id] = {"filename": filename, "chunks": ["失败残片"]}
                first_started.set()
                if not release_first.wait(3):
                    raise AssertionError("未释放并发测试门闩")
                raise BusinessException("第一份重建模拟失败", code=502)
            self.vectors.index_document(document_id, filename, chunks)

        def delete(document_id):
            trace.append((local.name, "delete"))
            self.vectors.delete_document(document_id)

        def rebuild(name):
            local.name = name
            with sessions() as db:
                if name == "second":
                    # 模拟等锁前已读过旧 ORM 对象；锁后必须重新取最新状态。
                    db.get(Document, document_id)
                    second_attempted.set()
                try:
                    return document_service.reindex_document(db, document_id)
                finally:
                    if name == "second":
                        second_finished.set()

        with patch.object(document_service.rag_service, "index_document", index), \
             patch.object(document_service.rag_service, "delete_document", delete), \
             ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(rebuild, "first")
            try:
                self.assertTrue(first_started.wait(2))
                second = pool.submit(rebuild, "second")
                self.assertTrue(second_attempted.wait(2))
                self.assertFalse(second_finished.wait(0.1), "第二份重建必须等待第一份补偿完成")
                self.assertEqual(trace, [("first", "index")])
                release_first.set()
                with self.assertRaises(BusinessException):
                    first.result(timeout=3)
                self.assertEqual(second.result(timeout=3).status, "indexed")
            finally:
                release_first.set()
        self.assertEqual(trace, [("first", "index"), ("first", "delete"), ("second", "index")])
        self.assertIn(document_id, self.vectors.rows)
        self.assertNotEqual(self.vectors.rows[document_id]["chunks"], ["失败残片"])
        with sessions() as db:
            self.assertEqual(db.get(Document, document_id).status, "indexed")

    def test_delete_waits_for_reindex_failure_cleanup_then_removes_document(self):
        sessions, document_id = self.concurrent_document()
        first_started = threading.Event()
        delete_attempted = threading.Event()
        delete_finished = threading.Event()
        release_first = threading.Event()
        local = threading.local()
        trace = []

        def index(document_id, filename, chunks):
            trace.append("reindex")
            self.vectors.rows[document_id] = {"filename": filename, "chunks": ["失败残片"]}
            first_started.set()
            if not release_first.wait(3):
                raise AssertionError("未释放并发测试门闩")
            raise BusinessException("重建模拟失败", code=502)

        def delete(document_id):
            trace.append(local.name)
            self.vectors.delete_document(document_id)

        def rebuild():
            local.name = "reindex-compensation"
            with sessions() as db:
                return document_service.reindex_document(db, document_id)

        def remove():
            local.name = "delete"
            with sessions() as db:
                db.get(Document, document_id)
                delete_attempted.set()
                try:
                    document_service.delete_document(db, document_id)
                finally:
                    delete_finished.set()

        with patch.object(document_service.rag_service, "index_document", index), \
             patch.object(document_service.rag_service, "delete_document", delete), \
             ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(rebuild)
            try:
                self.assertTrue(first_started.wait(2))
                deletion = pool.submit(remove)
                self.assertTrue(delete_attempted.wait(2))
                self.assertFalse(delete_finished.wait(0.1), "删除必须等待重建补偿完成")
                self.assertEqual(trace, ["reindex"])
                release_first.set()
                with self.assertRaises(BusinessException):
                    first.result(timeout=3)
                deletion.result(timeout=3)
            finally:
                release_first.set()
        self.assertEqual(trace, ["reindex", "reindex-compensation", "delete"])
        self.assertNotIn(document_id, self.vectors.rows)
        with sessions() as db:
            self.assertIsNone(db.get(Document, document_id))
        self.assertEqual(self.files(), [])

    def test_reindex_waiting_for_delete_reloads_and_refuses_deleted_row(self):
        sessions, document_id = self.concurrent_document()
        delete_started = threading.Event()
        rebuild_attempted = threading.Event()
        rebuild_finished = threading.Event()
        release_delete = threading.Event()
        index_calls = []

        def delete(document_id):
            delete_started.set()
            if not release_delete.wait(3):
                raise AssertionError("未释放并发测试门闩")
            self.vectors.delete_document(document_id)

        def index(document_id, filename, chunks):
            index_calls.append(document_id)
            self.vectors.index_document(document_id, filename, chunks)

        def remove():
            with sessions() as db:
                document_service.delete_document(db, document_id)

        def rebuild():
            with sessions() as db:
                # 保留引用，确保 identity map 中确实有等锁前读取的旧对象。
                old_row = db.get(Document, document_id)
                self.assertIsNotNone(old_row)
                rebuild_attempted.set()
                try:
                    return document_service.reindex_document(db, document_id)
                finally:
                    rebuild_finished.set()

        with patch.object(document_service.rag_service, "index_document", index), \
             patch.object(document_service.rag_service, "delete_document", delete), \
             ThreadPoolExecutor(max_workers=2) as pool:
            deletion = pool.submit(remove)
            try:
                self.assertTrue(delete_started.wait(2))
                rebuilding = pool.submit(rebuild)
                self.assertTrue(rebuild_attempted.wait(2))
                self.assertFalse(rebuild_finished.wait(0.1))
                release_delete.set()
                deletion.result(timeout=3)
                with self.assertRaises(BusinessException) as raised:
                    rebuilding.result(timeout=3)
                self.assertEqual(raised.exception.code, 404)
            finally:
                release_delete.set()
        self.assertEqual(index_calls, [])
        self.assertNotIn(document_id, self.vectors.rows)
        with sessions() as db:
            self.assertIsNone(db.get(Document, document_id))
        self.assertEqual(self.files(), [])


if __name__ == "__main__":
    unittest.main()
