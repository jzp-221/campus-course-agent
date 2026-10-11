"""文档抽取与分块离线测试：内存文件，无网络、数据库或真实配置。"""

from io import BytesIO
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.common.exceptions import BusinessException
from app.services.document_text_service import chunk_text, extract_text


def pdf_bytes(*, text=None, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text is not None:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        contents = DecodedStreamObject()
        contents.set_data(f"BT /F1 12 Tf 20 700 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(contents)
    if encrypted:
        writer.encrypt("test-document-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(*, populated=True):
    document = Document()
    if populated:
        document.add_paragraph("课程规则")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "出勤"
        table.cell(0, 1).text = "不得低于 80%"
        document.add_paragraph("请按时提交作业")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class DocumentTextTests(unittest.TestCase):
    def assert_business_error(self, name, content, expected):
        with self.assertRaises(BusinessException) as raised:
            extract_text(name, content)
        self.assertEqual(raised.exception.code, 400)
        self.assertIn(expected, raised.exception.message)

    def test_txt_utf8_bom_and_normalized_paragraphs(self):
        data = "\ufeff  课程规则\r\n\r\n\r\n  出勤\t不得低于 80%。  \r结束\x00".encode("utf-8")
        self.assertEqual(extract_text("课程.TXT", data), "课程规则\n\n出勤 不得低于 80%。\n结束")

    def test_markdown_preserves_readable_structure(self):
        data = "# 课程规则\n\n- 按时上课\n- 作业独立完成".encode("utf-8")
        self.assertEqual(extract_text("rules.md", data), data.decode("utf-8"))

    def test_unsupported_extension_is_rejected(self):
        self.assert_business_error("legacy.doc", b"content", "仅支持")

    def test_invalid_utf8_is_rejected(self):
        self.assert_business_error("rules.txt", b"\xff\xfeinvalid", "UTF-8")

    def test_empty_file_is_rejected(self):
        self.assert_business_error("rules.txt", b"", "文件为空")

    def test_blank_text_is_rejected(self):
        self.assert_business_error("rules.md", b" \n\t\x00", "有效文本")

    def test_pdf_real_text_layer_is_extracted(self):
        self.assertEqual(
            extract_text("rules.pdf", pdf_bytes(text="Attendance minimum: 80 percent.")),
            "Attendance minimum: 80 percent.",
        )

    def test_pdf_without_text_layer_has_explicit_scan_message(self):
        self.assert_business_error("scan.pdf", pdf_bytes(), "扫描版 PDF")
        self.assert_business_error("scan.pdf", pdf_bytes(), "不支持 OCR")

    def test_encrypted_pdf_is_explicitly_rejected(self):
        self.assert_business_error("rules.pdf", pdf_bytes(encrypted=True), "已加密")

    def test_corrupt_pdf_is_rejected(self):
        self.assert_business_error("broken.pdf", b"%PDF-1.7\nbroken", "损坏或格式不正确")

    def test_docx_paragraph_and_table_order_is_preserved(self):
        self.assertEqual(
            extract_text("rules.docx", docx_bytes()),
            "课程规则\n出勤 不得低于 80%\n请按时提交作业",
        )

    def test_blank_docx_is_rejected(self):
        self.assert_business_error("blank.docx", docx_bytes(populated=False), "有效文本")

    def test_corrupt_docx_is_rejected_without_internal_error(self):
        self.assert_business_error("broken.docx", b"invalid zip", "损坏或格式不正确")
        with patch("docx.Document", side_effect=RuntimeError("/private/secret-storage/source.docx")):
            with self.assertRaises(BusinessException) as raised:
                extract_text("rules.docx", b"opaque content")
        self.assertNotIn("/private", raised.exception.message)
        self.assertNotIn("secret-storage", raised.exception.message)

    def test_chunks_have_exact_overlap_and_restore_all_characters(self):
        text = "".join(chr(0x4E00 + i) for i in range(1234))
        chunks = chunk_text(text)
        self.assertEqual([len(chunk) for chunk in chunks], [500, 500, 334])
        self.assertTrue(all(left[-50:] == right[:50] for left, right in zip(chunks, chunks[1:])))
        self.assertEqual(chunks[0] + "".join(chunk[50:] for chunk in chunks[1:]), text)

    def test_chunk_boundaries_do_not_produce_redundant_tail(self):
        self.assertEqual([len(chunk) for chunk in chunk_text("课" * 500)], [500])
        self.assertEqual([len(chunk) for chunk in chunk_text("课" * 501)], [500, 51])
        self.assertEqual([len(chunk) for chunk in chunk_text("课" * 950)], [500, 500])

    def test_empty_chunk_input_is_explicit_error(self):
        with self.assertRaises(BusinessException) as raised:
            chunk_text("\n\t  ")
        self.assertEqual(raised.exception.code, 400)
        self.assertIn("为空", raised.exception.message)

    def test_short_sections_do_not_mix_topics(self):
        text = "# 课程规定\n课程编码：TEST101\n\n## 作业\n迟交扣 2 分。\n\n### 考试\n不允许携带计算器。"
        chunks = chunk_text(text)
        self.assertEqual(len(chunks), 3)
        self.assertEqual(chunks[0], "# 课程规定\n课程编码：TEST101")
        self.assertEqual(chunks[1], "## 作业\n迟交扣 2 分。")
        self.assertEqual(chunks[2], "### 考试\n不允许携带计算器。")
        self.assertNotIn("计算器", chunks[1])
        self.assertNotIn("迟交", chunks[2])

    def test_every_window_keeps_document_and_chapter_prefix(self):
        section = "## 平时作业要求\n" + "".join(chr(0x4E00 + i) for i in range(1100))
        chunks = chunk_text(section + "\n\n## 考试规则\n禁止计算器。", filename="高等数学.md")
        prefix = "高等数学.md · 平时作业要求\n"
        self.assertTrue(all(chunk.startswith(prefix) for chunk in chunks[:-1]))
        bodies = [chunk[len(prefix):] for chunk in chunks[:-1]]
        self.assertEqual([len(body) for body in bodies], [500, 500, len(section) - 900])
        self.assertEqual(bodies[0] + "".join(body[50:] for body in bodies[1:]), section)
        self.assertTrue(all(left[-50:] == right[:50] for left, right in zip(bodies, bodies[1:])))
        self.assertEqual(chunks[-1], "高等数学.md · 考试规则\n## 考试规则\n禁止计算器。")

    def test_plain_text_has_source_prefix_without_changing_window_size(self):
        chunks = chunk_text("课" * 501, filename="课程规定.txt")
        prefix = "课程规定.txt · 正文\n"
        self.assertTrue(all(chunk.startswith(prefix) for chunk in chunks))
        self.assertEqual([len(chunk.removeprefix(prefix)) for chunk in chunks], [500, 51])

    def test_fourth_level_heading_does_not_split_a_chapter(self):
        text = "## 考试规则\n考试闭卷。\n#### 注意事项\n不允许计算器。\n## 作业规则\n必须独立完成。"
        chunks = chunk_text(text, filename="课程.md")
        self.assertEqual(len(chunks), 2)
        self.assertIn("#### 注意事项", chunks[0])
        self.assertNotIn("作业规则", chunks[0])

    def test_demo_math_splits_into_introduction_and_six_chapters(self):
        from pathlib import Path

        source = Path(__file__).resolve().parents[1] / "docs/demo/课程资料/高等数学学习与考核说明.md"
        chunks = chunk_text(source.read_text(encoding="utf-8"), filename=source.name)
        self.assertEqual(len(chunks), 7)
        calculator_chunks = [chunk for chunk in chunks if "计算器" in chunk]
        self.assertEqual(len(calculator_chunks), 1)
        self.assertTrue(calculator_chunks[0].startswith(source.name + " · 五、考试允许携带的物品\n"))
        self.assertNotIn("平时作业要求", calculator_chunks[0])


if __name__ == "__main__":
    unittest.main()
